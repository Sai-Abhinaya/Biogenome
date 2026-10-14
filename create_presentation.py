import os
import math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ==============================================================================
# COLOR PALETTE SPECIFICATION (Academic & Scientific)
# ==============================================================================
NAVY_DARK    = RGBColor(11, 25, 44)       # Deep Navy #0B192C (Header/Hero)
NAVY_CARD    = RGBColor(19, 39, 67)       # Navy Card fill
TEAL         = RGBColor(10, 147, 150)     # Primary Accent Teal #0A9396
TEAL_LIGHT   = RGBColor(230, 246, 246)    # Light Teal tint for badges
MINT_GREEN   = RGBColor(45, 106, 79)      # Muted Green #2D6A4F
MINT_LIGHT   = RGBColor(235, 246, 240)    # Soft Mint background
GREEN_ACCENT = RGBColor(82, 183, 136)     # Vibrant Sage/Mint #52B788
BG_LIGHT     = RGBColor(245, 247, 250)    # Off-White / Light Slate background
WHITE        = RGBColor(255, 255, 255)    # Pure White
TEXT_DARK    = RGBColor(15, 23, 42)       # Slate 900 for high-contrast body
TEXT_MUTED   = RGBColor(100, 116, 139)    # Slate 500 for secondary text
TEXT_LIGHT   = RGBColor(241, 245, 249)    # Slate 100 for dark backgrounds
BORDER_GRAY  = RGBColor(226, 232, 240)    # Slate 200 border stroke
CARD_BG      = RGBColor(255, 255, 255)    # Card fill
TAG_BG       = RGBColor(241, 245, 249)    # Badge background

FONT_HEADING = "Segoe UI"
FONT_BODY    = "Segoe UI"

def set_shape_flat_fill(shape, fill_color, border_color=None, border_width=1):
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(border_width)
    else:
        shape.line.fill.background()

def add_header(slide, title_text, category_text, subtitle_text):
    """Adds a standard scientific header to content slides."""
    # Top banner bar
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
    set_shape_flat_fill(top_bar, NAVY_DARK)
    
    # Category Pill
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.18), Inches(2.8), Inches(0.28))
    set_shape_flat_fill(pill, TEAL)
    tf_p = pill.text_frame
    tf_p.word_wrap = True
    tf_p.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_p = tf_p.paragraphs[0]
    p_p.text = category_text.upper()
    p_p.font.name = FONT_HEADING
    p_p.font.size = Pt(9.5)
    p_p.font.bold = True
    p_p.font.color.rgb = WHITE
    p_p.alignment = PP_ALIGN.CENTER
    
    # Title Text
    tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.48), Inches(11.7), Inches(0.42))
    tf_t = tb_title.text_frame
    tf_t.word_wrap = True
    tf_t.margin_top = Inches(0)
    tf_t.margin_bottom = Inches(0)
    tf_t.margin_left = Inches(0)
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.name = FONT_HEADING
    p_t.font.size = Pt(20)
    p_t.font.bold = True
    p_t.font.color.rgb = WHITE
    
    # Subtitle Text
    tb_sub = slide.shapes.add_textbox(Inches(0.8), Inches(0.88), Inches(11.7), Inches(0.24))
    tf_s = tb_sub.text_frame
    tf_s.word_wrap = True
    tf_s.margin_top = Inches(0)
    tf_s.margin_bottom = Inches(0)
    tf_s.margin_left = Inches(0)
    p_s = tf_s.paragraphs[0]
    p_s.text = subtitle_text
    p_s.font.name = FONT_BODY
    p_s.font.size = Pt(11)
    p_s.font.color.rgb = GREEN_ACCENT

def add_footer(slide, current_slide, total_slides=13):
    """Adds a standard footer to content slides."""
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.02))
    set_shape_flat_fill(line, BORDER_GRAY)
    
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(7.12), Inches(8.0), Inches(0.25))
    tf = tb.text_frame
    tf.margin_left = Inches(0)
    p = tf.paragraphs[0]
    p.text = "BioGenome — DNA Sequence Analysis Using Data Structures & Algorithms (DSA-3)"
    p.font.name = FONT_BODY
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED
    
    tb_num = slide.shapes.add_textbox(Inches(10.5), Inches(7.12), Inches(2.0), Inches(0.25))
    tf_n = tb_num.text_frame
    tf_n.margin_right = Inches(0)
    p_n = tf_n.paragraphs[0]
    p_n.text = f"Slide {current_slide} of {total_slides}"
    p_n.font.name = FONT_BODY
    p_n.font.size = Pt(9.5)
    p_n.font.bold = True
    p_n.font.color.rgb = TEAL
    p_n.alignment = PP_ALIGN.RIGHT

def create_card(slide, left, top, width, height, title, bg_color=WHITE, border_color=BORDER_GRAY):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    set_shape_flat_fill(card, bg_color, border_color, border_width=1.2)
    return card

# Initialize Presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_slide_layout = prs.slide_layouts[6]

print("Initialized PowerPoint Presentation in 16:9 layout.")
